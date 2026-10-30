from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "matlab_stage3_4" / "220GHz太赫兹阶段三四融合复现汇报_适中篇幅版.docx"
FIGURE_ROOT = ROOT / "outputs" / "matlab_stage3_4" / "figures"

FIGURES = [
    ("fig01_measured_pattern_radius_scales.png", "图1 实测方向图与两种扫描半径尺度", "H/V 半功率波束宽度约为 0.7026°/0.8328°，论文半径却为 2°–5°。因此原 LED 扫描尺度不能直接用于 220 GHz。图中的 0.0703°–0.3513°是按较窄 HPBW 的 0.1–0.5 倍设定的本地 THz 工程范围，不是论文参数。"),
    ("fig02_stage3_radius_scheduler_covariance.png", "图2 双向调度与自适应半径", "ns=8 时，两端每 8 拍交换 Active/Passive，半径互斥，形成门控脉冲。论文策略触及 2°下限，THz 策略则保持亚度尺度。该图验证调度和半径逻辑，不代表动态性能达标。"),
    ("fig03_stage3_static_pointing_estimates.png", "图3 静态场景真实误差与估计", "图中同时比较真实指向误差和滤波估计。原 MATLAB 骨架在 Gp=0.98 时有估计曲线贴近零的启动瞬态，因此不能单看估计值判断是否对准，还要看真实误差和控制方向。"),
    ("fig04_stage3_passive_invariants_psd.png", "图4 Passive 冻结与协方差诊断", "Passive 端只更新 q 与 P(1,1)，角度、角速度和其他协方差项原样携带。160 拍静态测试的受保护项变化为 0，通过 1e-12 断言。Schur 补下界用于守住协方差半正定性，不会增加角度观测。"),
    ("fig05_stage4_dynamic_mode_comparison.png", "图5 正交移动动态误差、距离与半径", "速度因子 0.2、dt=0.5 s、200 步，距离从 2.3000 m 变为 2.3087 m。误差、半径和 LOS 角速度需合看：动态问题不只是扫描半径，还涉及预测时标和控制生效时刻。"),
    ("fig06_stage4_filter_control_link_diagnostics.png", "图6 动态滤波、控制与链路诊断", "估计、控制量、功率、创新和协方差用于定位失锁环节。当前链路是简化模型，功率不能当成校准接收机的实测 dBm/SNR。"),
    ("fig07_stage4_axis_error_components.png", "图7 两端分轴误差", "总误差会掩盖轴间差异。前馈可能改善一个轴、同时恶化另一个轴，因此需要核对 omega 的定义、两端局部坐标符号和控制限幅，不能把两轴平均成总体改善。"),
    ("fig08_stage4_orthogonal_trajectory.png", "图8 Fig. 9 正交运动轨迹", "Rover 水平移动、Elevator 垂直移动，轨迹对应 Fig. 9 构型。轨迹正确只说明运动输入符合设定，不等同于闭环指向收敛。"),
    ("fig09_author_fig8_distance_archive.png", "图9 论文 Fig. 8 多距离归档重绘", "作者 MAT 数据覆盖 0.92–5.06 m：EKF 误差由 0.899°升至 5.166°，ES 在 0.92 m 为 28.250°。这是论文归档重绘，不是本地 THz 实验。"),
    ("fig10_author_fig10_speed_archive.png", "图10 论文 Fig. 10 速度归档重绘", "作者归档中速度因子增至 1.3 时，EKF 误差达到 19.511°。扫描期间运动增大后，准静态假设和局部线性化逐渐失效；该图用于对照论文趋势，不是本地动态结果。"),
]


def set_font(run, name="宋体", size=10.5, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Cm(0.74)
    return p


def figure(doc, item):
    filename, caption, note = item
    figure_number = caption.split(" ", 1)[0]
    path = FIGURE_ROOT / filename
    if not path.is_file():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Inches(6.05))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(caption), size=9, bold=True)
    body(doc, f"{figure_number}解读：{note}")


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(4)
    for name, size in [("Heading 1", 14), ("Heading 2", 11.5)]:
        style = doc.styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run("220 GHz 太赫兹双向主动对准：阶段三、四复现汇报"), "黑体", 17, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run("实现过程与仿真结果"), "楷体", 10.5)

    doc.add_heading("完成内容与论文依据", level=1)
    body(doc, "本轮从手写 MATLAB 一、二阶段骨架继续推进。原代码用 LED 高斯角度响应和 PI 反馈验证滤波、控制符号及静态闭环；迁移到 220 GHz 后，我换入实测 H/V 天线方向图，并把状态扩成 x=[q,q_dot,e_H,e_E,omega_H,omega_E]^T。统一仿真入口现已包含双向时分、Passive 安全更新、两种半径和正交移动测试。")
    body(doc, "阶段三沿用论文 III-D 的 Active/Passive 交替和式(27)–(31)标量更新：Passive 时只更新 q、P(1,1)，角度、角速度及相关协方差项冻结。为避免固定交叉项造成协方差非半正定，本地另加 Schur 补下界；这是数值保护，不是论文新增观测。半径方面保留式(25)、(26)的角协方差最大特征值和 2°–5°论文限幅，并按作者 NodeEKF 的窗口均值尺度换算 K_delta=1/(7000·ns)。")
    body(doc, "THz 适配部分是本地改动：实测 HPBW 只有约 0.7026°/0.8328°，因此另设 0.0703°–0.3513°半径，取较窄 HPBW 的 0.1–0.5 倍。六状态及异步样本年龄参考师兄 Python 接口；前馈按角速度乘预测时延加入，本地 tau_ff=8 s 是长时隙假设，不是 Python 默认机械延迟。方向图与传播损耗已接入，但噪声、双端耦合和坐标变换仍有简化。")

    doc.add_heading("阶段三：静态交替验证", level=1)
    body(doc, "静态结果见图2–图4：速度设为 0，运行 160 拍，R1/R2 各 Active 80 拍、Passive 80 拍。两种半径下，受保护状态及协方差项最大变化都是 0，通过 1e-12 断言。这说明被动期不会因无信息的六维更新而产生角状态漂移；它验证的是实现不变量，不是动态跟踪性能。")
    for item in FIGURES[:4]:
        figure(doc, item)

    doc.add_heading("阶段四：正交移动与结果", level=1)
    body(doc, "动态结果见图5–图8。测试按 Fig. 9 构型运行：速度因子 0.2、dt=0.5 s、200 步，距离由 2.3000 m 变到 2.3087 m。四组固定种子单次结果全部未收敛：论文半径无前馈 Rover/Elevator 为 31.017°/22.361°，有前馈为 17.031°/30.640°；THz 半径无前馈为 20.632°/19.386°，有前馈为 25.794°/14.492°。")
    body(doc, "结合图5的总误差、图6的估计与控制诊断、图7的分轴误差和图8的运动轨迹来看，前馈对两轴的影响不一致，THz 亚度半径也没有单独带来收敛。我的判断是先验证 omega 的定义、两端局部坐标符号、参考时刻和命令限幅，再调 PI；否则单次曲线无法区分时序问题和控制方向问题。")
    for item in FIGURES[4:8]:
        figure(doc, item)

    doc.add_heading("论文归档对照与结论", level=1)
    body(doc, "图9和图10来自论文作者 MAT 归档，不是本地 THz 实验。图9的多距离数据中，EKF 误差由 0.92 m 时的 0.899°升到 5.06 m 时的 5.166°，ES 在 0.92 m 为 28.250°；图10显示速度因子到 1.3 时 EKF 误差达到 19.511°。这与距离增加导致量测变弱、运动加快后准静态假设变差的解释一致。")
    for item in FIGURES[8:]:
        figure(doc, item)
    body(doc, "总体上，阶段三的 Passive 冻结防线已通过严格检查，阶段四的 THz 动态闭环尚未收敛。论文半径远宽于实测主瓣，THz 半径更贴近天线尺度但仍不能解决长时标和符号接口问题。旧 LED 模型的 0.8155°/1.1977°是历史基线，不是本轮 THz 成绩。下一步先做零速前馈等价和 H/E 正负角速度注入测试，再开展速度扫描与多随机种子验证。")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("220 GHz 太赫兹阶段三、四复现 | 本地仿真与论文归档分别标注"), size=9)
    doc.core_properties.title = "220 GHz 太赫兹阶段三、四融合复现汇报"
    doc.core_properties.subject = "实现方法与仿真结果"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
